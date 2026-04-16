from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass


TABLE_START = "\n<!-- TABLE_START -->\n"
TABLE_END = "\n<!-- TABLE_END -->\n"
LIST_START = "\n<!-- LIST_START -->\n"
LIST_END = "\n<!-- LIST_END -->\n"
LIST_BULLET_MARKERS = (
    "\u2022",
    "\u25cf",
    "\u25cb",
    "\u25e6",
    "\u25aa",
    "\u25b8",
    "\u25ba",
    "\u25a0",
    "\u25a1",
    "\u25c6",
    "\u25c7",
    "-",
    "*",
    "\u2013",
    "\u2014",
)


@dataclass(slots=True)
class ParsedSegmentRecord:
    segment_type: str
    segment_number: int | None
    title: str | None
    text: str


@dataclass(slots=True)
class ParsedSourceDocument:
    text: str
    segments: list[ParsedSegmentRecord]
    extraction_method: str
    requires_fallback_ocr: bool = False


@dataclass(slots=True)
class ChunkDraft:
    chunk_index: int
    text: str
    normalized_text: str
    token_count: int
    contains_structure: bool
    dedup_hash: str
    start_page: int | None = None
    end_page: int | None = None
    start_slide: int | None = None
    end_slide: int | None = None
    heading_path: list[str] | None = None
    tags: dict | None = None


def _normalize_list_line(line: str) -> str:
    roman = re.match(r"^(\s*)([ivxlcdmIVXLCDM]+)[\.\)]\s*(.*)$", line)
    if roman:
        indent, marker, content = roman.groups()
        return f"{indent}{marker}. {content}"
    numbered = re.match(r"^(\s*)(\d+)[\.\)]\s*(.*)$", line)
    if numbered:
        indent, marker, content = numbered.groups()
        return f"{indent}{marker}. {content}"
    lettered = re.match(r"^(\s*)([a-zA-Z])[\.\)]\s*(.*)$", line)
    if lettered:
        indent, marker, content = lettered.groups()
        return f"{indent}{marker}. {content}"
    bulleted = re.match(r"^(\s*)[•●○◦▪▸►■□◆◇\\-*–—]\s*(.*)$", line)
    if bulleted:
        indent, content = bulleted.groups()
        return f"{indent}• {content}"
    return line


def _is_list_line(line: str) -> bool:
    patterns = [
        r"^[\s]*[•●○◦▪▸►■□◆◇]\s*",
        r"^[\s]*\d+[\.\)]\s*",
        r"^[\s]*[a-zA-Z][\.\)]\s*",
        r"^[\s]*[ivxlcdmIVXLCDM]+[\.\)]\s*",
        r"^[\s]*[-*–—]\s+",
    ]
    return any(re.match(pattern, line) for pattern in patterns)


def _normalize_list_line_safe(line: str) -> str:
    roman = re.match(r"^(\s*)([ivxlcdmIVXLCDM]+)[\.\)]\s*(.*)$", line)
    if roman:
        indent, marker, content = roman.groups()
        return f"{indent}{marker}. {content}"
    numbered = re.match(r"^(\s*)(\d+)[\.\)]\s*(.*)$", line)
    if numbered:
        indent, marker, content = numbered.groups()
        return f"{indent}{marker}. {content}"
    lettered = re.match(r"^(\s*)([a-zA-Z])[\.\)]\s*(.*)$", line)
    if lettered:
        indent, marker, content = lettered.groups()
        return f"{indent}{marker}. {content}"
    bulleted = re.match(r"^(\s*)(.*)$", line)
    if bulleted:
        indent, content = bulleted.groups()
        if content and content[0] in LIST_BULLET_MARKERS:
            return f"{indent}\u2022 {content[1:].lstrip()}"
    return line


def _is_list_line_safe(line: str) -> bool:
    stripped = line.lstrip()
    if not stripped:
        return False
    if stripped[0] in LIST_BULLET_MARKERS:
        return True
    patterns = [
        r"^\d+[\.\)]\s*",
        r"^[a-zA-Z][\.\)]\s*",
        r"^[ivxlcdmIVXLCDM]+[\.\)]\s*",
    ]
    return any(re.match(pattern, stripped) for pattern in patterns)


def _group_list_blocks(text: str) -> str:
    lines = text.splitlines()
    result: list[str] = []
    in_list = False
    list_buffer: list[str] = []
    for raw_line in lines:
        line = _normalize_list_line_safe(raw_line.rstrip())
        if _is_list_line_safe(line):
            if not in_list:
                in_list = True
                list_buffer = []
            list_buffer.append(line)
            continue
        if in_list:
            result.append(LIST_START)
            result.extend(list_buffer)
            result.append(LIST_END)
            list_buffer = []
            in_list = False
        result.append(line)
    if in_list and list_buffer:
        result.append(LIST_START)
        result.extend(list_buffer)
        result.append(LIST_END)
    return "\n".join(result).strip()


def _table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    max_cols = max(len(row) for row in rows)
    normalized_rows: list[list[str]] = []
    for row in rows:
        cleaned = [str(cell or "").strip().replace("\n", " ").replace("|", "\\|") for cell in row]
        while len(cleaned) < max_cols:
            cleaned.append("")
        normalized_rows.append(cleaned)
    output = ["| " + " | ".join(normalized_rows[0]) + " |", "| " + " | ".join(["---"] * max_cols) + " |"]
    output.extend("| " + " | ".join(row) + " |" for row in normalized_rows[1:])
    return "\n".join(output)


def _sanitize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return _group_list_blocks(text).strip()


class StructurePreservingParser:
    def parse(self, raw_bytes: bytes, filename: str) -> ParsedSourceDocument:
        extension = filename.lower().rsplit(".", maxsplit=1)[-1] if "." in filename else ""
        if extension == "docx":
            return self._parse_docx(raw_bytes, filename)
        if extension == "pdf":
            return self._parse_pdf(raw_bytes, filename)
        if extension == "pptx":
            return self._parse_pptx(raw_bytes, filename)
        return self._parse_text(raw_bytes, filename)

    def _parse_text(self, raw_bytes: bytes, filename: str) -> ParsedSourceDocument:
        text = _sanitize_text(raw_bytes.decode("utf-8", errors="replace"))
        return ParsedSourceDocument(
            text=text,
            segments=[ParsedSegmentRecord(segment_type="section", segment_number=1, title=filename, text=text)],
            extraction_method="plain_text",
        )

    def _parse_docx(self, raw_bytes: bytes, filename: str) -> ParsedSourceDocument:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover - runtime dependency
            raise RuntimeError("python-docx is required for DOCX parsing") from exc

        doc = DocxDocument(io.BytesIO(raw_bytes))
        content_parts: list[str] = []
        for element in doc.element.body:
            if element.tag.endswith("tbl"):
                for table in doc.tables:
                    if table._tbl is element:
                        rows = [[cell.text for cell in row.cells] for row in table.rows]
                        table_md = _table_to_markdown(rows)
                        if table_md:
                            content_parts.append(f"{TABLE_START}{table_md}{TABLE_END}")
                        break
            elif element.tag.endswith("p"):
                for para in doc.paragraphs:
                    if para._p is element:
                        text = para.text.strip()
                        if text:
                            content_parts.append(_normalize_list_line(text))
                        break
        content = _sanitize_text("\n".join(content_parts))
        return ParsedSourceDocument(
            text=content,
            segments=[ParsedSegmentRecord(segment_type="section", segment_number=1, title=filename, text=content)],
            extraction_method="python-docx",
        )

    def _parse_pdf(self, raw_bytes: bytes, filename: str) -> ParsedSourceDocument:
        parser_errors: list[str] = []
        for parser_name, parser in (
            ("pdfplumber", self._parse_pdf_with_pdfplumber),
            ("pymupdf", self._parse_pdf_with_pymupdf),
            ("pypdf", self._parse_pdf_with_pypdf),
        ):
            try:
                segments = parser(raw_bytes)
                if segments:
                    text = _sanitize_text("\n\n".join(segment.text for segment in segments))
                    density = len(text.strip()) / max(len(raw_bytes), 1)
                    return ParsedSourceDocument(
                        text=text,
                        segments=segments,
                        extraction_method=parser_name,
                        requires_fallback_ocr=density < 0.03,
                    )
            except Exception as exc:  # pragma: no cover - parser fallback
                parser_errors.append(f"{parser_name}: {exc}")
        raise RuntimeError("Unable to parse PDF with local extractors: " + "; ".join(parser_errors))

    def _parse_pdf_with_pdfplumber(self, raw_bytes: bytes) -> list[ParsedSegmentRecord]:
        import pdfplumber

        segments: list[ParsedSegmentRecord] = []
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                parts: list[str] = []
                for table in page.extract_tables() or []:
                    if table:
                        table_md = _table_to_markdown([[cell or "" for cell in row] for row in table])
                        if table_md:
                            parts.append(f"{TABLE_START}{table_md}{TABLE_END}")
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
                if parts:
                    segments.append(
                        ParsedSegmentRecord(
                            segment_type="page",
                            segment_number=page_num,
                            title=f"Page {page_num}",
                            text=_sanitize_text("\n\n".join(parts)),
                        )
                    )
        return segments

    def _parse_pdf_with_pymupdf(self, raw_bytes: bytes) -> list[ParsedSegmentRecord]:
        import fitz

        segments: list[ParsedSegmentRecord] = []
        with fitz.open(stream=raw_bytes, filetype="pdf") as doc:
            for page_num, page in enumerate(doc, start=1):
                page_text = page.get_text("text") or ""
                if page_text.strip():
                    segments.append(
                        ParsedSegmentRecord(
                            segment_type="page",
                            segment_number=page_num,
                            title=f"Page {page_num}",
                            text=_sanitize_text(page_text),
                        )
                    )
        return segments

    def _parse_pdf_with_pypdf(self, raw_bytes: bytes) -> list[ParsedSegmentRecord]:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw_bytes))
        segments: list[ParsedSegmentRecord] = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                segments.append(
                    ParsedSegmentRecord(
                        segment_type="page",
                        segment_number=page_num,
                        title=f"Page {page_num}",
                        text=_sanitize_text(page_text),
                    )
                )
        return segments

    def _parse_pptx(self, raw_bytes: bytes, filename: str) -> ParsedSourceDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - runtime dependency
            raise RuntimeError("python-pptx is required for PPTX parsing") from exc

        presentation = Presentation(io.BytesIO(raw_bytes))
        segments: list[ParsedSegmentRecord] = []
        for slide_num, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = getattr(shape, "text", "").strip()
                    if text:
                        parts.append(text)
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    table_md = _table_to_markdown(rows)
                    if table_md:
                        parts.append(f"{TABLE_START}{table_md}{TABLE_END}")
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                notes = ""
            if notes:
                parts.append(f"Speaker Notes:\n{notes}")
            if parts:
                segments.append(
                    ParsedSegmentRecord(
                        segment_type="slide",
                        segment_number=slide_num,
                        title=f"Slide {slide_num}",
                        text=_sanitize_text("\n\n".join(parts)),
                    )
                )
        content = _sanitize_text("\n\n".join(segment.text for segment in segments))
        return ParsedSourceDocument(
            text=content,
            segments=segments or [ParsedSegmentRecord(segment_type="slide", segment_number=1, title=filename, text="")],
            extraction_method="python-pptx",
        )


class StructureAwareChunker:
    def __init__(self, *, chunk_size: int = 2000, chunk_overlap: int = 200) -> None:
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def split(self, document: ParsedSourceDocument) -> list[ChunkDraft]:
        protected_regions = self._identify_protected_regions(document.text)
        segments = self._segment_text(document.text, protected_regions)
        chunks: list[str] = []
        current = ""
        contains_structure = False
        structure_flags: list[bool] = []
        for segment in segments:
            if segment["protected"]:
                if current and len(current) + len(segment["text"]) > self.chunk_size * 1.5:
                    chunks.append(current.strip())
                    structure_flags.append(contains_structure)
                    current = ""
                    contains_structure = False
                current += segment["text"]
                contains_structure = True
                continue
            for sub_chunk in self._split_regular_text(segment["text"]):
                if len(current) + len(sub_chunk) <= self.chunk_size:
                    current += sub_chunk
                else:
                    if current.strip():
                        chunks.append(current.strip())
                        structure_flags.append(contains_structure)
                    current = sub_chunk
                    contains_structure = False
        if current.strip():
            chunks.append(current.strip())
            structure_flags.append(contains_structure)

        draft_chunks: list[ChunkDraft] = []
        for index, chunk_text in enumerate(chunks, start=1):
            cleaned = chunk_text.replace(TABLE_START, "\n").replace(TABLE_END, "\n").replace(LIST_START, "\n").replace(LIST_END, "\n")
            cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
            token_count = max(len(cleaned.split()), 1)
            start_page = self._first_segment_number(document.segments, "page", cleaned)
            end_page = self._last_segment_number(document.segments, "page", cleaned)
            start_slide = self._first_segment_number(document.segments, "slide", cleaned)
            end_slide = self._last_segment_number(document.segments, "slide", cleaned)
            normalized = cleaned.lower()
            draft_chunks.append(
                ChunkDraft(
                    chunk_index=index,
                    text=cleaned,
                    normalized_text=normalized,
                    token_count=token_count,
                    contains_structure=structure_flags[index - 1],
                    dedup_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
                    start_page=start_page,
                    end_page=end_page,
                    start_slide=start_slide,
                    end_slide=end_slide,
                    heading_path=[document.segments[0].title] if document.segments and document.segments[0].title else None,
                    tags={"contains_structure": structure_flags[index - 1], "extraction_method": document.extraction_method},
                )
            )
        return draft_chunks

    def _identify_protected_regions(self, text: str) -> list[tuple[int, int]]:
        regions: list[tuple[int, int]] = []
        for start_marker, end_marker in ((TABLE_START, TABLE_END), (LIST_START, LIST_END)):
            for match in re.finditer(re.escape(start_marker) + r"(.*?)" + re.escape(end_marker), text, re.DOTALL):
                regions.append((match.start(), match.end()))
        return sorted(regions, key=lambda region: region[0])

    def _segment_text(self, text: str, protected_regions: list[tuple[int, int]]) -> list[dict]:
        if not protected_regions:
            return [{"text": text, "protected": False}]
        segments: list[dict] = []
        cursor = 0
        for start, end in protected_regions:
            if cursor < start:
                segments.append({"text": text[cursor:start], "protected": False})
            segments.append({"text": text[start:end], "protected": True})
            cursor = end
        if cursor < len(text):
            segments.append({"text": text[cursor:], "protected": False})
        return [segment for segment in segments if segment["text"].strip()]

    def _split_regular_text(self, text: str) -> list[str]:
        if len(text) <= self.chunk_size:
            return [text]
        chunks: list[str] = []
        current = ""
        for paragraph in text.split("\n\n"):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            if len(current) + len(paragraph) + 2 <= self.chunk_size:
                current += paragraph + "\n\n"
                continue
            if current.strip():
                chunks.append(current.strip())
            if len(paragraph) <= self.chunk_size:
                current = paragraph + "\n\n"
                continue
            sentence_chunk = ""
            for sentence in re.split(r"(?<=[.!?])\s+", paragraph):
                if len(sentence_chunk) + len(sentence) + 1 <= self.chunk_size:
                    sentence_chunk += sentence + " "
                    continue
                if sentence_chunk.strip():
                    chunks.append(sentence_chunk.strip())
                sentence_chunk = sentence + " "
            current = sentence_chunk.strip()
        if current.strip():
            chunks.append(current.strip())
        return chunks

    def _first_segment_number(self, segments: list[ParsedSegmentRecord], segment_type: str, chunk_text: str) -> int | None:
        for segment in segments:
            if segment.segment_type == segment_type and segment.text and segment.text[:120] in chunk_text:
                return segment.segment_number
        return None

    def _last_segment_number(self, segments: list[ParsedSegmentRecord], segment_type: str, chunk_text: str) -> int | None:
        for segment in reversed(segments):
            if segment.segment_type == segment_type and segment.text and segment.text[:120] in chunk_text:
                return segment.segment_number
        return None
