from __future__ import annotations

import io
import re
from pathlib import Path

from app.services.types import NormalizedDocument, NormalizedSegment

PAGE_OR_SLIDE_RE = re.compile(r"^(page|slide)\s+(\d+)\b[:\-\s]*", re.IGNORECASE)
MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
LIST_ITEM_RE = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+|[ivxlcdm]+\.\s+)", re.IGNORECASE)


def _flush_block(
    segments: list[NormalizedSegment],
    lines: list[str],
    *,
    structure_kind: str,
    segment_type: str,
    segment_number: int | None,
    active_heading: str | None,
    heading_path: list[str],
) -> None:
    if not lines:
        return
    text = "\n".join(lines).strip()
    if not text:
        return
    segments.append(
        NormalizedSegment(
            segment_type=segment_type,
            segment_number=segment_number,
            title=active_heading,
            text=text,
            metadata={
                "structure_kind": structure_kind,
                "heading_path": heading_path.copy(),
                "contains_structure": structure_kind != "paragraph",
            },
        )
    )


def _normalize_heading_path(heading_path: list[str], level: int, title: str) -> list[str]:
    trimmed = heading_path[: max(level - 1, 0)]
    trimmed.append(title)
    return trimmed


def _stable_table_markdown(rows: list[list[str | None]]) -> str:
    cleaned_rows = [[(cell or "").strip() for cell in row] for row in rows if any((cell or "").strip() for cell in row)]
    if not cleaned_rows:
        return ""
    width = max(len(row) for row in cleaned_rows)
    normalized = [row + [""] * (width - len(row)) for row in cleaned_rows]
    header = normalized[0]
    divider = ["---"] * width
    body = normalized[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(divider) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _segments_from_text(text: str) -> list[NormalizedSegment]:
    segments: list[NormalizedSegment] = []
    lines = text.splitlines()
    block: list[str] = []
    heading_path: list[str] = []
    active_heading: str | None = None
    segment_number: int | None = None
    structure_kind = "paragraph"

    for raw_line in lines:
        line = raw_line.rstrip()
        if not line.strip():
            _flush_block(
                segments,
                block,
                structure_kind=structure_kind,
                segment_type="content",
                segment_number=segment_number,
                active_heading=active_heading,
                heading_path=heading_path,
            )
            block = []
            structure_kind = "paragraph"
            continue

        page_match = PAGE_OR_SLIDE_RE.match(line.strip())
        if page_match:
            _flush_block(
                segments,
                block,
                structure_kind=structure_kind,
                segment_type="content",
                segment_number=segment_number,
                active_heading=active_heading,
                heading_path=heading_path,
            )
            block = []
            segment_number = int(page_match.group(2))
            continue

        heading_match = MARKDOWN_HEADING_RE.match(line.strip())
        if heading_match:
            _flush_block(
                segments,
                block,
                structure_kind=structure_kind,
                segment_type="content",
                segment_number=segment_number,
                active_heading=active_heading,
                heading_path=heading_path,
            )
            block = []
            active_heading = heading_match.group(2).strip()
            heading_path = _normalize_heading_path(heading_path, len(heading_match.group(1)), active_heading)
            segments.append(
                NormalizedSegment(
                    segment_type="heading",
                    segment_number=segment_number,
                    title=active_heading,
                    text=active_heading,
                    metadata={"structure_kind": "heading", "heading_path": heading_path.copy(), "contains_structure": True},
                )
            )
            structure_kind = "paragraph"
            continue

        if line.lstrip().startswith("|"):
            structure_kind = "table"
        elif LIST_ITEM_RE.match(line):
            structure_kind = "list"
        block.append(line)

    _flush_block(
        segments,
        block,
        structure_kind=structure_kind,
        segment_type="content",
        segment_number=segment_number,
        active_heading=active_heading,
        heading_path=heading_path,
    )
    return segments


def _pdf_with_pdfplumber(raw_bytes: bytes) -> tuple[str, str]:
    import pdfplumber

    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            lines.append(f"Page {index}")
            page_text = page.extract_text(layout=True) or ""
            if page_text.strip():
                lines.append(page_text.strip())
            for table in page.extract_tables() or []:
                table_text = _stable_table_markdown(table)
                if table_text:
                    lines.append(table_text)
    return "\n\n".join(lines).strip(), "pdfplumber"


def _pdf_with_pymupdf(raw_bytes: bytes) -> tuple[str, str]:
    import fitz

    lines: list[str] = []
    with fitz.open(stream=raw_bytes, filetype="pdf") as document:
        for index, page in enumerate(document, start=1):
            lines.append(f"Page {index}")
            text = page.get_text("text").strip()
            if text:
                lines.append(text)
    return "\n\n".join(lines).strip(), "pymupdf"


def _pdf_with_pypdf(raw_bytes: bytes) -> tuple[str, str]:
    from pypdf import PdfReader

    lines: list[str] = []
    reader = PdfReader(io.BytesIO(raw_bytes))
    for index, page in enumerate(reader.pages, start=1):
        lines.append(f"Page {index}")
        text = (page.extract_text() or "").strip()
        if text:
            lines.append(text)
    return "\n\n".join(lines).strip(), "pypdf"


def _parse_pdf(raw_bytes: bytes) -> tuple[str, str]:
    parsers = (_pdf_with_pdfplumber, _pdf_with_pymupdf, _pdf_with_pypdf)
    last_error: Exception | None = None
    for parser in parsers:
        try:
            text, method = parser(raw_bytes)
        except Exception as exc:  # pragma: no cover - parser fallback path
            last_error = exc
            continue
        if text.strip():
            return text, method
    if last_error is not None:
        raise ValueError(f"PDF parsing failed: {last_error}") from last_error
    raise ValueError("PDF parsing produced no text")


def _parse_docx(raw_bytes: bytes) -> tuple[str, str]:
    from docx import Document

    document = Document(io.BytesIO(raw_bytes))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                lines.append(f"# {text}")
            else:
                lines.append(text)
    for table in document.tables:
        table_text = _stable_table_markdown([[cell.text for cell in row.cells] for row in table.rows])
        if table_text:
            lines.append(table_text)
    return "\n\n".join(lines).strip(), "python-docx"


def _parse_pptx(raw_bytes: bytes) -> tuple[str, str]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw_bytes))
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"Slide {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                table_text = _stable_table_markdown(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if table_text:
                    lines.append(table_text)
        if getattr(slide, "notes_slide", None) is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
    return "\n\n".join(lines).strip(), "python-pptx"


def parse_source_bytes(raw_bytes: bytes, filename: str) -> NormalizedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text, method = _parse_pdf(raw_bytes)
    elif suffix == ".docx":
        text, method = _parse_docx(raw_bytes)
    elif suffix == ".pptx":
        text, method = _parse_pptx(raw_bytes)
    elif suffix in {".md", ".markdown", ".txt"}:
        text, method = raw_bytes.decode("utf-8", errors="replace"), "plain-text"
    else:
        raise ValueError(f"Unsupported source format: {suffix or filename}")
    segments = _segments_from_text(text)
    return NormalizedDocument(
        text=text.strip(),
        segments=segments,
        metadata={"extraction_method": method, "requires_fallback_ocr": suffix == ".pdf" and not text.strip()},
    )
